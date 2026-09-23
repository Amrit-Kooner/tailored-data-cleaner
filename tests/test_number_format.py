"""
Regression tests for number handling in the Excel output:

  * ONLY money columns get the 2-decimal ".00" format.
  * Every other number keeps its own value and decimals (10 stays 10, 44.3
    stays 44.3, 2024 has no thousands comma).
  * Measurement columns (height, weight, quantity ...) become real numbers.
  * Price-name detection is whole-word, not "contains the letters".

Run from the project root:
    python -m unittest discover -s tests -v
"""

import os
import random
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from uni_cleaner.detectors import looks_like_price, looks_like_measure
from uni_cleaner.loaders import fixed_decimals_from_format
from uni_cleaner.output.excel import save_as_formatted_excel
from uni_cleaner.pipeline import clean_file

TWO_DP = "0.00"   # prices: 2 decimals, never a thousands comma


def build_xlsx(header, rows):
    path = os.path.join(tempfile.mkdtemp(), "in.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def run_to_excel(header, rows):
    """Clean an xlsx and write it exactly as batch.save_outputs does. Returns (result, Cleaned worksheet)."""
    res = clean_file(build_xlsx(header, rows))
    out = os.path.join(tempfile.mkdtemp(), "o.xlsx")
    save_as_formatted_excel(res, out, "in.xlsx")
    return res, openpyxl.load_workbook(out)["Cleaned"]


def cell(ws, col_name, row=2):
    headers = [c.value for c in ws[1]]
    return ws.cell(row, headers.index(col_name) + 1)


class OnlyMoneyGetsTwoDecimals(unittest.TestCase):
    def setUp(self):
        random.seed(3)
        header = ["part", "height", "weight", "quantity", "pack_size", "year", "rating", "unit_price", "score"]
        rows = [[f"P{i:03d}",
                 random.choice([10, 12.5, 15]),
                 round(random.uniform(1, 90), 2),
                 random.choice([1, 5, 10]),
                 random.choice([6, 12, 24]),
                 random.choice([2024, 2025, 2026]),
                 random.choice([1, 2, 3, 4, 5]),
                 random.choice([4, 10.5, 44.3]),
                 random.choice([0.5, 1.5])] for i in range(60)]
        self.res, self.ws = run_to_excel(header, rows)

    def test_price_column_has_two_decimals_even_for_whole_numbers(self):
        self.assertEqual(cell(self.ws, "unit_price").number_format, TWO_DP)
        self.assertIn("unit_price", self.res.money_cols)

    def test_no_other_column_has_the_two_decimal_format(self):
        for name in ("height", "weight", "quantity", "pack_size", "year", "rating", "score"):
            with self.subTest(column=name):
                self.assertEqual(cell(self.ws, name).number_format, "General")

    def test_year_has_no_thousands_comma_or_decimals(self):
        c = cell(self.ws, "year")
        self.assertIn(c.value, (2024, 2025, 2026))
        self.assertEqual(c.number_format, "General")

    def test_only_price_is_in_money_cols(self):
        self.assertEqual(self.res.money_cols, ["unit_price"])


class DecimalsAreNeverTruncated(unittest.TestCase):
    def test_44_30_stays_44_3_not_44(self):
        _, ws = run_to_excel(["item", "weight", "unit_price"],
                             [["a", 44.30, 44.30], ["b", 10, 10], ["c", 3.125, 3.125], ["d", 7.5, 7.5]])
        self.assertEqual(cell(ws, "weight", 2).value, 44.3)      # not 44
        self.assertEqual(cell(ws, "weight", 4).value, 3.125)     # not rounded to 3.13
        self.assertEqual(cell(ws, "unit_price", 2).value, 44.3)
        self.assertEqual(cell(ws, "unit_price", 2).number_format, TWO_DP)   # displays 44.30

    def test_measure_values_survive_exactly(self):
        vals = [51.66, 12.5, 100, 0.001, 44.375]
        res, _ = run_to_excel(["item", "weight"], [[f"i{i}", v] for i, v in enumerate(vals)])
        self.assertEqual(res.df["weight"].tolist(), vals)


class NoThousandsCommaAnywhere(unittest.TestCase):
    def test_big_numbers_in_every_kind_of_column_have_no_comma(self):
        header = ["part", "amount", "length", "dimension", "quantity"]
        rows = [[f"P{i}", 2365.53 + i, 1400 + i * 305, 203000 + i * 1000, 10000 + i] for i in range(12)]
        res, ws = run_to_excel(header, rows)
        for name in ("amount", "length", "dimension", "quantity"):
            with self.subTest(column=name):
                self.assertNotIn(",", cell(ws, name).number_format)
        self.assertEqual(cell(ws, "amount").number_format, TWO_DP)          # price: 2365.53
        for name in ("length", "dimension", "quantity"):                    # non-price: plain
            self.assertEqual(cell(ws, name).number_format, "General")


class RawDisplayIsPreserved(unittest.TestCase):
    """'if the raw data already has .00 keep it; a plain 100 must not become 100.00'"""

    def _xlsx_with_formats(self):
        path = os.path.join(tempfile.mkdtemp(), "fmt.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["part", "weight", "length"])
        data = [("a", 4.3, 100), ("b", 100, 1400.5), ("c", 5.25, 7), ("d", 0.35, 200), ("e", 12, 55)]
        for row in data:
            ws.append(list(row))
        ws["B2"].number_format = "0.00"       # 4.3   shows 4.30
        ws["B3"].number_format = "0.00"       # 100   shows 100.00  (raw already had .00)
        ws["C3"].number_format = "#,##0.00"   # 1400.5 shows 1,400.50 -> keep 2 decimals, drop the comma
        wb.save(path)
        return path

    def _write(self, path):
        res = clean_file(path)
        out = os.path.join(tempfile.mkdtemp(), "o.xlsx")
        save_as_formatted_excel(res, out, os.path.basename(path))
        return res, openpyxl.load_workbook(out)["Cleaned"]

    def test_excel_cell_shown_with_two_decimals_keeps_them(self):
        res, ws = self._write(self._xlsx_with_formats())
        self.assertEqual(cell(ws, "weight", 2).value, 4.3)
        self.assertEqual(cell(ws, "weight", 2).number_format, "0.00")      # 4.30, not 4
        self.assertEqual(cell(ws, "weight", 3).number_format, "0.00")      # raw showed 100.00
        self.assertEqual(cell(ws, "length", 3).number_format, "0.00")      # comma dropped, decimals kept

    def test_general_cells_stay_general(self):
        res, ws = self._write(self._xlsx_with_formats())
        self.assertEqual(cell(ws, "weight", 4).number_format, "General")   # 5.25
        self.assertEqual(cell(ws, "weight", 6).number_format, "General")   # 12   -> not 12.00
        self.assertEqual(cell(ws, "length", 2).number_format, "General")   # 100  -> not 100.00
        self.assertEqual(cell(ws, "length", 6).number_format, "General")

    def test_values_are_never_changed_by_preserving_the_display(self):
        res, ws = self._write(self._xlsx_with_formats())
        self.assertEqual([cell(ws, "weight", r).value for r in range(2, 7)], [4.3, 100, 5.25, 0.35, 12])

    def test_text_cells_typed_with_trailing_zeros_keep_them(self):
        header, rows = ["part", "weight"], [["a", "4.30"], ["b", "100"], ["c", "0.35"], ["d", "12.50"]]
        res, ws = run_to_excel(header, rows)
        self.assertEqual(cell(ws, "weight", 2).value, 4.3)
        self.assertEqual(cell(ws, "weight", 2).number_format, "0.00")      # "4.30" stays 4.30
        self.assertEqual(cell(ws, "weight", 3).number_format, "General")   # "100" stays 100
        self.assertEqual(cell(ws, "weight", 4).number_format, "General")   # "0.35"
        self.assertEqual(cell(ws, "weight", 5).number_format, "0.00")      # "12.50" stays 12.50

    def test_csv_text_keeps_trailing_zeros_too(self):
        path = os.path.join(tempfile.mkdtemp(), "w.csv")
        with open(path, "w", encoding="utf-8") as f:
            f.write("part,weight\na,4.30\nb,100\nc,0.35\nd,12.50\n")
        res = clean_file(path)
        self.assertEqual(res.number_display["weight"], {4.3: 2, 12.5: 2})

    def test_price_columns_ignore_raw_display_and_always_get_two_decimals(self):
        res, ws = run_to_excel(["part", "unit_price"], [["a", 4.5], ["b", 100], ["c", 2365.53], ["d", 7]])
        for r in range(2, 6):
            self.assertEqual(cell(ws, "unit_price", r).number_format, TWO_DP)


class FormatParser(unittest.TestCase):
    def test_fixed_decimals_from_excel_format(self):
        cases = {"0.00": 2, "#,##0.00": 2, '"GBP"#,##0.000': 3, "[$\u00a3-809]#,##0.00": 2,
                 "#,##0.00_);(#,##0.00)": 2, "0.0": 1,
                 "General": None, "0": None, "#,##0": None, "0.0#": None, "0.00%": None,
                 "dd/mm/yyyy": None, "0.00E+00": None, "@": None, "": None}
        for fmt, expected in cases.items():
            with self.subTest(fmt=fmt):
                self.assertEqual(fixed_decimals_from_format(fmt), expected)


class MeasurementColumns(unittest.TestCase):
    def test_measure_columns_become_real_numbers(self):
        rows = [[f"P{i}", 10 + i * 1.5, 3 + i] for i in range(12)]
        res, ws = run_to_excel(["part", "height", "weight"], rows)
        for name in ("height", "weight"):
            self.assertTrue(str(res.df[name].dtype).startswith(("float", "int")), name)
            self.assertIsInstance(cell(ws, name).value, (int, float))

    def test_mostly_unique_measure_is_not_mistaken_for_a_text_id(self):
        rows = [[f"P{i}", round(1.1 * i + 0.37, 2)] for i in range(30)]   # 100% unique
        res, _ = run_to_excel(["part", "weight"], rows)
        self.assertTrue(str(res.df["weight"].dtype).startswith("float"))

    def test_dimension_text_is_left_alone(self):
        rows = [[f"P{i}", "10x20", ] for i in range(6)]
        res, _ = run_to_excel(["part", "dimension"], rows)
        self.assertEqual(set(res.df["dimension"]), {"10x20"})
        self.assertFalse(any(s["step"] == "dtype_measure_numeric" for s in res.log["steps"]))

    def test_unparseable_measure_values_are_kept_and_flagged_red(self):
        rows = [["a", 10], ["b", 12], ["c", "about 5"], ["d", 8]]
        res, _ = run_to_excel(["part", "height"], rows)
        self.assertIn("about 5", res.df["height"].astype(str).tolist())
        self.assertEqual(int(res.fail_mask["height"].sum()), 1)

    def test_total_weight_and_total_qty_are_numbers_not_prices(self):
        rows = [[f"P{i}", 5 + i * 0.5, 10 + i] for i in range(10)]
        res, ws = run_to_excel(["part", "total_weight", "total_qty"], rows)
        self.assertEqual(res.money_cols, [])
        self.assertEqual(cell(ws, "total_weight").number_format, "General")
        self.assertEqual(cell(ws, "total_qty").number_format, "General")


class CurrencySymbolsCountAsMoney(unittest.TestCase):
    def test_currency_column_not_named_price_still_gets_two_decimals(self):
        rows = [["a", "£10"], ["b", "£4.5"], ["c", "£7"]]
        res, ws = run_to_excel(["item", "revenue"], rows)
        self.assertIn("revenue", res.money_cols)
        self.assertEqual(cell(ws, "revenue").number_format, TWO_DP)


class PriceNameMatchingIsWholeWord(unittest.TestCase):
    MONEY = ["price", "unit_price", "part_price", "list_price", "unitprice", "ListPrice", "netprice",
             "total_cost", "totalcost", "subtotal", "net_amount", "gross_value", "fee", "surcharge",
             "charge", "amount", "cost", "total", "day_rate", "prices"]
    NOT_MONEY = ["total_qty", "total_quantity", "total_weight", "total_items", "cost_centre",
                 "cost_code", "exchange_rate", "heart_rate", "rate_code", "value_type",
                 "attribute_value", "discharge_date", "pricelist", "priceband", "amount_of_items",
                 "charge_point_id", "fee_type", "cost_type", "total_rows", "rating", "valued_by",
                 "pricing_notes", "separate", "coffee", "customer_name",
                 "charge_on_finance_charge_flag", "price_flag", "price_flags", "priceflag",
                 "unit_price_flagged", "chargeflags", "Amount Flag", "flag_price"]

    def test_money_names_match(self):
        for n in self.MONEY:
            with self.subTest(name=n):
                self.assertTrue(looks_like_price(n))

    def test_non_money_names_do_not_match(self):
        for n in self.NOT_MONEY:
            with self.subTest(name=n):
                self.assertFalse(looks_like_price(n))

    def test_a_price_is_never_a_measure(self):
        for n in ("unit_price", "price", "total_cost"):
            self.assertFalse(looks_like_measure(n))


class FlagColumnsAreNeverPrices(unittest.TestCase):
    """A column with 'flag' in its name is a Y/N-style flag even if it also says 'price'/'charge'."""

    def setUp(self):
        rows = [["a", 0, 0, 5.0], ["b", 1, 1, 7.5], ["c", 0, 0, 2.0]]
        self.res, self.ws = run_to_excel(
            ["item", "price_flag", "charge_on_finance_charge_flag", "unit_price"], rows)

    def test_flag_columns_are_not_money_cols(self):
        self.assertNotIn("price_flag", self.res.money_cols)
        self.assertNotIn("charge_on_finance_charge_flag", self.res.money_cols)

    def test_real_price_column_still_is(self):
        self.assertIn("unit_price", self.res.money_cols)

    def test_flag_columns_get_no_two_decimal_format(self):
        self.assertNotEqual(cell(self.ws, "price_flag").number_format, TWO_DP)
        self.assertNotEqual(cell(self.ws, "charge_on_finance_charge_flag").number_format, TWO_DP)

    def test_zero_in_a_flag_column_is_not_flagged_as_a_zero_price(self):
        for col in ("price_flag", "charge_on_finance_charge_flag"):
            self.assertFalse(self.res.review_mask[col].any(), col)


if __name__ == "__main__":
    unittest.main(verbosity=2)
