"""
Regression tests for BLUE flagging of copies of the same field ("email" / "email.1"):

  A similar NAME never makes two columns BLUE. They are flagged only when the data
  in them is the same in every row, IGNORING CASE ("Acme" / "acme" / "ACME" are the
  same value) and blank only where the other is blank. Before, a copy that merely
  matched wherever both cells were filled (e.g. one column half empty) was flagged
  as well, because detect_duplicate_content_columns re-tested it with a looser rule.

  Duplicate ROWS are also matched ignoring case: "Acme" and "acme" in otherwise
  identical rows -> the later row is deleted.

Run from the project root:
    python -m unittest discover -s tests -v
"""

import os
import sys
import tempfile
import textwrap
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from uni_cleaner.pipeline import clean_file
from uni_cleaner.output.excel import save_as_formatted_excel, HEADER_REVIEW_FILL

BLUE_RGB = HEADER_REVIEW_FILL.start_color.rgb[-6:]   # "D9EAF7"


def run(csv_text):
    path = os.path.join(tempfile.mkdtemp(), "t.csv")
    with open(path, "w", encoding="utf-8") as f:
        f.write(textwrap.dedent(csv_text).strip() + "\n")
    return clean_file(path)


def header_fills(res):
    """
    {column name: fill rgb or None} for the header row of the "Review" sheet
    -- duplicate-content/unnamed-header flagging is column-level (not tied to
    specific rows), so it's shown on that sheet's headers regardless of which
    rows appear below. A text block may sit above the header row now, so the
    header row is located by its "Cleaned Row #" marker rather than assumed
    to be row 1.
    """
    out = os.path.join(tempfile.mkdtemp(), "out.xlsx")
    save_as_formatted_excel(res, out, "t.csv")
    ws = openpyxl.load_workbook(out)["Review"]
    header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Cleaned Row #")
    return {c.value: (c.fill.start_color.rgb[-6:] if c.fill.fill_type else None)
            for c in ws[header_row] if c.column > 1}


class SameFieldCopiesAreBlueOnlyWhenTheDataIsTheSame(unittest.TestCase):
    def test_identical_copy_is_blue(self):
        r = run("""
            id,email,email
            1,a@x.com,a@x.com
            2,b@x.com,b@x.com
        """)
        self.assertEqual(list(r.df.columns), ["id", "email", "email.1"])
        self.assertEqual(r.duplicate_name_map, {"email": ["email.1"], "email.1": ["email"]})
        fills = header_fills(r)
        self.assertEqual(fills["email"], BLUE_RGB)
        self.assertEqual(fills["email.1"], BLUE_RGB)
        self.assertIsNone(fills["id"])

    def test_different_copy_is_not_blue(self):
        r = run("""
            id,email,email
            1,a@x.com,z@x.com
            2,b@x.com,y@x.com
        """)
        self.assertEqual(r.duplicate_name_map, {})
        fills = header_fills(r)
        self.assertIsNone(fills["email"])
        self.assertIsNone(fills["email_alt_1"])

    def test_copy_with_blanks_where_the_other_has_data_is_not_blue(self):
        # the ONE row where both are filled matches -- but the rest is blank vs data
        r = run("""
            id,email,email
            1,a@x.com,a@x.com
            2,b@x.com,
            3,c@x.com,
        """)
        self.assertEqual(list(r.df.columns), ["id", "email", "email_alt_1"])
        self.assertEqual(r.duplicate_name_map, {})
        fills = header_fills(r)
        self.assertIsNone(fills["email"])
        self.assertIsNone(fills["email_alt_1"])

    def test_copy_with_one_differing_row_is_not_blue(self):
        r = run("""
            id,email,email
            1,a@x.com,a@x.com
            2,b@x.com,b@x.com
            3,c@x.com,other@x.com
        """)
        self.assertEqual(r.duplicate_name_map, {})

    def test_copies_that_differ_only_by_case_ARE_blue(self):
        # "notes" is not case-normalised by the pipeline, so this is a real case-only difference:
        # "Acme" / "acme" / "ACME" are the same value -> same data -> BLUE, and NOT renamed _alt_
        r = run("""
            id,notes,notes
            1,Acme,acme
            2,Beta,BETA
            3,Gamma,gAmMa
        """)
        self.assertEqual(list(r.df.columns), ["id", "notes", "notes.1"])
        self.assertEqual(r.df["notes"].tolist(), ["Acme", "Beta", "Gamma"])     # values untouched
        self.assertEqual(r.df["notes.1"].tolist(), ["acme", "BETA", "gAmMa"])
        self.assertEqual(r.duplicate_name_map, {"notes": ["notes.1"], "notes.1": ["notes"]})
        fills = header_fills(r)
        self.assertEqual(fills["notes"], BLUE_RGB)
        self.assertEqual(fills["notes.1"], BLUE_RGB)

    def test_case_only_difference_plus_one_really_different_row_is_not_blue(self):
        r = run("""
            id,notes,notes
            1,Acme,ACME
            2,Beta,other
        """)
        self.assertEqual(list(r.df.columns), ["id", "notes", "notes_alt_1"])
        self.assertEqual(r.duplicate_name_map, {})

    def test_case_only_difference_but_a_blank_against_a_value_is_not_blue(self):
        r = run("""
            id,notes,notes
            1,Acme,ACME
            2,Beta,
        """)
        self.assertEqual(r.duplicate_name_map, {})

    def test_two_blank_filled_copies_with_matching_blank_cells_are_still_blue(self):
        # blank in BOTH at the same rows is "the same" -- only a blank-vs-value clash is a difference
        r = run("""
            id,notes,notes
            1,Acme,Acme
            2,,
            3,Beta,Beta
        """)
        self.assertEqual(r.duplicate_name_map, {"notes": ["notes.1"], "notes.1": ["notes"]})

    def test_three_copies_only_the_exact_ones_are_blue(self):
        r = run("""
            id,email,email,email
            1,a@x.com,a@x.com,a@x.com
            2,b@x.com,b@x.com,
            3,c@x.com,c@x.com,
        """)
        self.assertEqual(list(r.df.columns), ["id", "email", "email.1", "email_alt_1"])
        self.assertEqual(r.duplicate_name_map, {"email": ["email.1"], "email.1": ["email"]})
        fills = header_fills(r)
        self.assertEqual(fills["email"], BLUE_RGB)
        self.assertEqual(fills["email.1"], BLUE_RGB)
        self.assertIsNone(fills["email_alt_1"])

    def test_all_copies_identical_are_all_blue_and_linked_to_each_other(self):
        r = run("""
            id,email,email,email
            1,a@x.com,a@x.com,a@x.com
            2,b@x.com,b@x.com,b@x.com
        """)
        self.assertEqual(sorted(r.duplicate_name_map["email.1"]), ["email", "email.2"])
        self.assertEqual(sorted(r.duplicate_name_map["email.2"]), ["email", "email.1"])

    def test_unrelated_similar_names_are_not_blue(self):
        # same prefix, same-looking names, different data
        r = run("""
            id,email,email_backup,email_verified
            1,a@x.com,z@x.com,yes
            2,b@x.com,y@x.com,no
        """)
        self.assertEqual(r.duplicate_name_map, {})


class DuplicateRowsAreMatchedIgnoringCase(unittest.TestCase):
    def test_rows_that_differ_only_by_case_are_deleted(self):
        r = run("""
            id,notes,owner
            1,Acme,Zed
            1,ACME,zed
            1,acme,ZED
            2,Beta,Yan
        """)
        self.assertEqual(len(r.df), 2)
        self.assertEqual(r.df["notes"].tolist(), ["Acme", "Beta"])     # first occurrence kept, as written
        self.assertEqual(r.df["owner"].tolist(), ["Zed", "Yan"])

    def test_rows_that_differ_in_more_than_case_are_kept(self):
        r = run("""
            id,notes,owner
            1,Acme,Zed
            1,ACME,Other
        """)
        self.assertEqual(len(r.df), 2)


class DifferentNamesKeepTheirOwnRule(unittest.TestCase):
    """Columns with different names are still matched on identical data where both are populated."""

    def test_different_names_same_data_is_blue(self):
        r = run("""
            id,contact,owner
            1,Acme,Acme
            2,Beta,Beta
        """)
        self.assertEqual(r.duplicate_name_map, {"contact": ["owner"], "owner": ["contact"]})

    def test_different_names_different_data_is_not_blue(self):
        r = run("""
            id,contact,owner
            1,Acme,Zed
            2,Beta,Yan
        """)
        self.assertEqual(r.duplicate_name_map, {})


if __name__ == "__main__":
    unittest.main()
