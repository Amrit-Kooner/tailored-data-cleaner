import os, sys, tempfile, textwrap, unittest
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from uni_cleaner.pipeline import clean_file
from uni_cleaner.output.excel import save_as_formatted_excel

class RequestedSheetNamesAndColours(unittest.TestCase):
    def run_clean(self, csv_text):
        path = os.path.join(tempfile.mkdtemp(), 't.csv')
        with open(path, 'w', encoding='utf-8') as f:
            f.write(textwrap.dedent(csv_text).strip() + '\n')
        return clean_file(path)

    def test_requested_sheet_names_and_distinct_colours(self):
        r = self.run_clean('''
            app,email,email
            A1,a@x.com,a@x.com
            A1,b@x.com,b@x.com
        ''')
        out = os.path.join(tempfile.mkdtemp(), 'out.xlsx')
        save_as_formatted_excel(r, out, 't.csv')
        wb = openpyxl.load_workbook(out)
        self.assertIn('ID Anomalies', wb.sheetnames)
        self.assertIn('ID Duplicates', wb.sheetnames)
        self.assertIn('Duplicate Values', wb.sheetnames)
        self.assertIn('Header Anomalies', wb.sheetnames)

        colours = {name: wb[name].sheet_properties.tabColor.rgb[-6:] for name in (
            'ID Anomalies', 'ID Duplicates', 'Duplicate Values', 'Header Anomalies'
        )}
        self.assertEqual(colours['ID Anomalies'], 'CDB4DB')
        self.assertEqual(colours['ID Duplicates'], 'F4B6C2')
        self.assertEqual(colours['Duplicate Values'], '92D050')
        self.assertEqual(colours['Header Anomalies'], 'D9EAF7')
        self.assertEqual(len(set(colours.values())), 4)

    def test_id_flags_remain_on_the_two_app_sheets(self):
        r = self.run_clean('''
            app,email
            A1,a@x.com
            A1,b@x.com
        ''')
        out = os.path.join(tempfile.mkdtemp(), 'out.xlsx')
        save_as_formatted_excel(r, out, 't.csv')
        wb = openpyxl.load_workbook(out)
        # ID Duplicates still has the duplicate ID rows.
        self.assertIn('A1', [c.value for row in wb['ID Duplicates'].iter_rows() for c in row])
        # The APP-ID anomaly sheet still exists for the separate invalid/blank-ID flagging path.
        self.assertIn('ID Anomalies', wb.sheetnames)

if __name__ == '__main__':
    unittest.main()
