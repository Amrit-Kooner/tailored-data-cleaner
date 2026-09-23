"""
Regression tests for APP-only uniqueness/anomaly handling and the split
issue sheets. Other identifiers are not treated as special APP fields.
"""
import os
import sys
import tempfile
import textwrap
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from uni_cleaner.pipeline import clean_file
from uni_cleaner.output.excel import (
    save_as_formatted_excel,
    APP_DUPLICATE_FILL,
    APP_INVALID_FILL,
    PRICE_FILL,
    DUPLICATE_VALUE_FILL,
)


def run(csv_text, name="t.csv"):
    path = os.path.join(tempfile.mkdtemp(), name)
    with open(path, "w", encoding="utf-8") as f:
        f.write(textwrap.dedent(csv_text).strip() + "\n")
    return clean_file(path)


class AppOnlyRules(unittest.TestCase):
    def test_duplicate_app_collapses_with_description(self):
        result = run("""
            app_number,description,price,other_id
            A1,short,10,B1
            A1,much longer description,20,B1
            A2,other,0,B2
        """)
        self.assertEqual(result.df["app_number"].tolist(), ["A1", "A2"])
        self.assertIn("A2", result.df["app_number"].tolist())

    def test_other_id_is_not_special_app_id(self):
        result = run("""
            app_number,description,other_id
            A1,one,B1
            A2,two,B1
        """)
        self.assertFalse(result.app_duplicate_mask["other_id"].any())
        self.assertFalse(result.app_invalid_mask["other_id"].any())
        self.assertNotIn("other_id", result.log["steps"][-1].get("step", ""))

    def test_split_issue_sheets_and_colours(self):
        result = run("""
            app_number,description,price,price_copy
            A1,one,10,10
            A2,two,0,0
            A2,three,-5,-5
            ,four,20,20
        """)
        out = os.path.join(tempfile.mkdtemp(), "out.xlsx")
        save_as_formatted_excel(result, out, "t.csv")
        wb = openpyxl.load_workbook(out)

        self.assertEqual(
            wb.sheetnames[:3],
            ["Cleaned", "Raw", "Report"]
        )
        self.assertIn("Duplicate Values", wb.sheetnames)
        self.assertIn("ID Duplicates", wb.sheetnames)
        self.assertIn("ID Anomalies", wb.sheetnames)
        self.assertIn("Price Anomalies", wb.sheetnames)
        self.assertNotIn("Review", wb.sheetnames)
        self.assertNotIn("APP Number Issues", wb.sheetnames)

        dup = wb["Duplicate Values"]
        app_invalid = wb["ID Anomalies"]
        price = wb["Price Anomalies"]

        # price and price_copy hold identical data -> pink header (row 2; row 1 is the note)
        self.assertEqual(dup["D2"].fill.fgColor.rgb[-6:], DUPLICATE_VALUE_FILL.start_color.rgb[-6:])
        # the A2/A2 pair gets collapsed by dedupe_by_app_number before the mask ever
        # runs (there's a description column), so the blank app_number is the only
        # row left to flag on the ID Anomalies sheet
        self.assertEqual(app_invalid["B2"].fill.fgColor.rgb[-6:], APP_INVALID_FILL.start_color.rgb[-6:])
        self.assertEqual(price["D2"].fill.fgColor.rgb[-6:], PRICE_FILL.start_color.rgb[-6:])

    def test_duplicate_app_sheet_when_no_description_to_collapse_by(self):
        """Without a description column, dedupe_by_app_number never runs, so a
        genuine duplicate APP value survives to be flagged on its own sheet."""
        result = run("""
            app_number,price
            A1,10
            A2,20
            A2,30
        """)
        out = os.path.join(tempfile.mkdtemp(), "out.xlsx")
        save_as_formatted_excel(result, out, "t.csv")
        wb = openpyxl.load_workbook(out)
        app_dupe = wb["ID Duplicates"]
        self.assertEqual(app_dupe["B2"].fill.fgColor.rgb[-6:], APP_DUPLICATE_FILL.start_color.rgb[-6:])


if __name__ == "__main__":
    unittest.main()
