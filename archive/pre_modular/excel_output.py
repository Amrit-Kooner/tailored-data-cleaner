# excel_output.py
"""
Takes a finished, cleaned dataframe and writes it out as a real, properly
formatted .xlsx.
"""

from datetime import datetime
import numbers
import pandas as pd
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


MISSING_FILL = PatternFill(start_color="FFD9A6", end_color="FFD9A6", fill_type="solid")  # orange
FAILED_FILL = PatternFill(start_color="FFB3B3", end_color="FFB3B3", fill_type="solid")   # red
REVIEW_FILL = PatternFill(start_color="A6C8FF", end_color="A6C8FF", fill_type="solid")   # blue

DUPLICATE_NAME_FILL = PatternFill(start_color="A6C8FF", end_color="A6C8FF", fill_type="solid")
DUPLICATE_NAME_FONT = Font(bold=True)

NUMBER_FORMAT_2DP = "#,##0.00"

# Accepted input formats for a pre-formatted date string. The first is what
# convert_to_uk_format() writes now (%Y = four-digit year); the second is
# retained so older output files continue to round-trip.
DATE_STRING_FORMATS = ("%d/%m/%Y", "%d/%m/%y")


def _is_real_number(value):
    if isinstance(value, bool):
        return False
    if value is None:
        return False
    if isinstance(value, numbers.Number):
        try:
            return not pd.isna(value)
        except (TypeError, ValueError):
            return True
    return False


def _parse_date_string(value):
    for fmt in DATE_STRING_FORMATS:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def save_as_formatted_excel(df, path, date_cols, missing_mask=None, fail_mask=None, review_mask=None,
                             duplicate_name_cols=None):
    wb = openpyxl.Workbook()
    ws = wb.active

    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        if duplicate_name_cols and col_name in duplicate_name_cols:
            cell.font = DUPLICATE_NAME_FONT
            cell.fill = DUPLICATE_NAME_FILL
            # BUG FIX: the "THINGS TO REVIEW" section of the audit report
            # tells the reviewer to look at the header cell's comment for
            # which column the duplicate matches. Previously we set the
            # font and fill but never the comment, so that instruction
            # pointed at an empty cell. Now the matching column name(s) are
            # written into a real Excel comment.
            matches = duplicate_name_cols.get(col_name) or []
            if matches:
                cell.comment = Comment(
                    "Data identical to: " + ", ".join(str(m) for m in matches),
                    "Cleaning Pipeline",
                )
        else:
            cell.font = Font(bold=True)

    for row_pos, (row_idx, row) in enumerate(zip(df.index, df.itertuples(index=False))):
        excel_row = row_pos + 2
        for col_idx, (col_name, value) in enumerate(zip(df.columns, row), start=1):
            is_date_col = col_name in date_cols
            is_missing = bool(missing_mask.at[row_idx, col_name]) if missing_mask is not None else False
            is_failed = bool(fail_mask.at[row_idx, col_name]) if fail_mask is not None else False
            is_review = bool(review_mask.at[row_idx, col_name]) if review_mask is not None else False

            if is_date_col and isinstance(value, str) and value.strip() and not is_failed:
                real_date = _parse_date_string(value)
                if real_date is not None:
                    cell = ws.cell(row=excel_row, column=col_idx, value=real_date)
                    cell.number_format = "DD/MM/YYYY"
                    continue
                # fall through and write as plain text if parsing fails

            excel_value = None if pd.isna(value) else value
            cell = ws.cell(row=excel_row, column=col_idx, value=excel_value)

            if _is_real_number(value):
                cell.number_format = NUMBER_FORMAT_2DP

            if is_failed:
                cell.fill = FAILED_FILL
            elif is_review:
                cell.fill = REVIEW_FILL
            elif is_missing:
                cell.fill = MISSING_FILL

    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = max(
            len(str(col_name)),
            df[col_name].astype(str).str.len().max() if len(df) > 0 else 0
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    wb.save(path)